from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path('/Users/zhangzaiyuan/Downloads/test')
OUT = ROOT / 'docs/浮屿作品说明与彩蛋指南.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.72), Inches(.65)
section.left_margin = section.right_margin = Inches(.8)
section.footer_distance = Inches(.28)

def set_font(style, name, size, color='263C35', bold=False):
    style.font.name, style.font.size = name, Pt(size)
    style.font.color.rgb, style.font.bold = RGBColor.from_string(color), bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)

set_font(doc.styles['Normal'], 'Heiti SC', 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(16.5)
normal.space_after = Pt(6)
set_font(doc.styles['Title'], 'Songti SC', 27, '000000')
doc.styles['Title'].paragraph_format.space_after = Pt(13)
set_font(doc.styles['Subtitle'], 'Heiti SC', 11, '60756A')
doc.styles['Subtitle'].paragraph_format.space_after = Pt(17)
set_font(doc.styles['Heading 1'], 'Songti SC', 19, '244D40', True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(11)
set_font(doc.styles['Heading 2'], 'Heiti SC', 12, '244D40', True)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)
for st in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    doc.styles[st].paragraph_format.widow_control = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('浮屿 ISOLA  ·  ')
r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('718078')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text):
    return doc.add_paragraph(text)

def sub(title, text):
    doc.add_heading(title, level=2)
    p(text)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i,w in enumerate(widths): t.columns[i].width = Inches(w)
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row): cells[i].text = text
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        keep = OxmlElement('w:cantSplit'); trpr.append(keep)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E3EDE4' if ri == 0 else ('F3F6F2' if ri % 2 else 'FFFFFF')); tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top',60),('bottom',60),('left',110),('right',110)]:
                el = OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(14.5)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5); run.font.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('浮屿作品说明与彩蛋指南', 'Title')
doc.add_paragraph('ISOLA  海岛观察手册  2026年10月6日', 'Subtitle')
p('浮屿是一座可以自由观察的三维海岛。三座山峰、沙滩、泻湖、山洞与码头构成了它的日常风景；日光、月相、四季和天气则不断改变同一处景物的表情。这份指南介绍容易错过的细节，也公开藏在视角按钮里的音乐彩蛋。')
sub('先把声音打开', '点击右上角的声音图标，开启背景音乐和自然声场。未开启时，按钮会用红色闪烁提醒。默认背景乐音量为80%，环境音为30%；可以在环境设置中分别调节。')
sub('靠近以后再看', '拖动可以围绕岛屿观察，滚轮或触控缩放可以拉近镜头。双击场景中的某个位置，镜头会靠近那个位置；山石、枝叶、亭子和岸边水面都值得停下来细看。底部的视角按钮可以直接带你去全岛、海岸、码头、浅海、泻湖、山洞、两座山亭、灯塔和赏月位置。')
table(['操作', '可以发现什么'], [
    ['漫游图标', '在声音图标后面。开启后，镜头缓缓环绕海岛。'],
    ['诗', '开启或关闭动态诗句。出现诗句还需要满足满月和画面可见条件。'],
    ['禅', '隐藏界面，只留下风景和符合条件的诗句；按钮变成“…”后可再次点击恢复。'],
    ['环境设置', '四季、天气、时间、月相、音量与画质都在同一个紧凑面板中。'],
], [1.25, 5.65])
sub('让变化自己发生', '默认开启昼夜轮换、四季轮换、天气自然变化和月相自动轮换，场景速率为0.5×。自动图标使用20%半透明高亮。月相按钮显示当前月亮形态，点击可依次选择八种月相，再进入自动模式；自动时带有小循环标记。')
p('进入画面前，加载文字会随机轮换“正在向海水加盐……”“正在呼叫鲸鱼和海豚……”等句子。海岛出现后，“张在元”印章先在画面中央盖下，再飞到左上角，与“浮屿”一起留下作品的落款。')

doc.add_page_break()
doc.add_heading('值得停下来观察的风景', level=1)
sub('太阳落下时的颜色', '晴朗白天能看到太阳与海面的日光。靠近日落，天空和云层逐渐染上金橙、玫红与更深的红色，太阳越接近海面，色调越暖。海面同时接住天空的颜色，远处缓慢移动的船影让夕阳更有“渔舟唱晚”的意味。')
sub('海上生明月', '夜里的月亮会缓缓升起、落下，月面使用真实纹理，带有暖金色的光。满月有更醒目的视觉尺度，其余月相采用较克制的比例。明月夜的岛屿仍保留可辨认的山石和植被，远处海天交界被薄雾柔化。')
sub('海水与海风', '海水会随光照和天气变化，在蓝色、青绿色与暮色之间转换。涌浪、近岸浪花和水面的细碎反光并不整齐重复；阵风还会改变浪的强弱。椰叶、树冠和草也会摆动，各处的摆幅与节奏有所不同。')
sub('云层和雨后的彩虹', '云朵缓缓移动，云量有时稀疏、有时密集。天气包含晴天、阴天、雨、雾、雷暴和雪。下过雨或雷暴后，如果转为有日光的晴天，天空会出现逐渐淡入又淡出的彩虹。想主动寻找它，可以先让雨落下来，再切到晴天等一等。')
sub('浅海和泻湖里的居民', '浅海有银蓝、黄纹和珊瑚色的鱼群，分布会随每次场景生成而改变。泻湖里则有常住的鱼群、海龟、螃蟹和龙虾：海龟摆动鳍足，螃蟹与龙虾会爬行、停顿，再继续活动。近景比全岛远景更容易看见这些动作。')
sub('远海里的偶遇', '鲸鱼在远海浮出、喷水，再重新潜下；海豚以数量和位置不同的小群出现，跃出水面后落回海中。两者错开出现。远处还有大小、船型、位置与航速各不相同的帆船和渔船轮廓，白天和夜晚都可能遇见。')
sub('山路尽头的停留点', '环岛路和石阶把海滩、山路与山顶连接起来，沿途有花草和小石子。灯塔所在山峰之外，另外两座山顶各有观景亭：木柱支撑，三面围合、三面长凳，入口与台阶相接。每座亭子的四个檐角都挂着小风铃，风吹时会摆动并发出铃声；冬日还能看到屋檐覆雪。')
sub('山洞的另一种气氛', '洞口被藤蔓和植被覆盖，前面是一片鲜艳的花朵。往里看，洞道蜿蜒进入山体，转角挡住了尽头。外面没有灯，深处只有一缕轻微起伏的暖光，让人隐约知道里面还有路。山顶灯塔则在入夜后点亮，扫光缓缓掠过周围。')

doc.add_page_break()
doc.add_heading('诗句与音乐怎样融入风景', level=1)
sub('只有看见满月才会遇见诗', '张若虚的《春江花月夜》只在满月、月亮实际出现在画面中且月光条件合适时展示。切到山洞、让月亮离开画面，或换成其他月相，诗句会暂时隐去。点击“赏月”可以进入适合看月亮的构图；白天点击时，场景会转到夜间。')
p('诗句使用有立体厚度的书法字形，在天空的空处出现，依次渐显、消散或飘离，正文去掉标点，署名保留“唐 张若虚”之间的空格。整首诗展示一遍后，这一轮月夜会收束到黎明，不再从第一联重播。')
p('禅模式保留诗句的动态效果。播放任何专属彩蛋歌曲时，诗会暂时隐藏；歌曲结束后，如果仍满足满月和画面可见条件，就恢复诗句的展示。')
sub('背景乐随着场景选择', '默认使用随景播放。背景曲以合成器乐营造气氛，专属彩蛋使用加入作品的歌曲录音。天气变化通常影响下一首曲子的选择；夕阳会转向《渔舟唱晚》的氛围。')
table(['场景', '常出现的背景曲'], [
    ['晴日', '彩云追月、山歌好比春江水、我的祖国'],
    ['明月夜', '春江花月夜、彩云追月、梁祝、女儿情、枉凝眉'],
    ['阴雨雾雪或静景', '平沙落雁、孟姜女哭长城、梁祝、枉凝眉'],
    ['夕阳与短暂余晖', '渔舟唱晚'],
], [1.65, 5.25])
sub('听见与动作对应的声音', '声场里有拍岸的海浪、回落的细碎浪花、风声、雨声和雷声，也有海鸟、亭子风铃、鲸鱼喷水、海豚与入水水花的声音。鲸鱼和海豚的声效跟随各自的出现事件，不必只盯着画面等待。')
sub('音乐速度和场景速度分开', '“随景音乐”右侧的三个小按钮分别控制播放或暂停、播放模式、音乐速度。播放模式在随景、单曲循环、顺序和随机之间切换；音乐速度有0.5×、0.75×、1×、1.25×、1.5×和2×。暂停背景乐时，环境声仍可继续播放。')
p('时间行的速率按钮控制昼夜和环境轮换，档位为0.5×、1×、2×、3×、4×、8×和16×。普通档位保留自然动作节奏；16×会把场景动画一起加速。在默认随景模式下，16×会随机固定一首背景曲循环，并减少密集音效，适合快速看一遍海岛的变化。')

doc.add_page_break()
doc.add_heading('音乐彩蛋指南', level=1)
p('彩蛋次数按底部的视角按钮累计。选定一个地点后，留在同一视角继续单击该按钮；切到其他视角会重新计数。建议让镜头到达后再继续点击，以便看清并听到触发结果。直接靠近山洞、码头或灯塔不会自动播放专属歌曲。')
table(['视角按钮', '累计点击', '触发歌曲'], [
    ['全岛', '第3次', '山歌好比春江水'],
    ['码头', '第3次', '送别'],
    ['南峰亭', '第3次', '功德宝山咒'],
    ['东峰亭', '第3次', '心经'],
    ['山洞', '第3次', '桔梗花'],
    ['山洞', '第5次', '波如来'],
    ['山洞', '第10次', '梦中的额吉'],
    ['山洞', '第12次', '大悲咒'],
    ['灯塔', '第3次', '月儿明风儿静'],
    ['灯塔', '第5次', '鲁冰花'],
    ['灯塔', '第10次', '花恋'],
    ['灯塔', '第12次', '驼糕之歌'],
], [1.55, 1.55, 3.8])
sub('歌曲只来一次', '每次触发的专属歌曲只播放一遍，结束后恢复普通背景音乐；停在原地不会反复循环同一首。山洞与灯塔可以继续累计到下一个指定次数，换到其他视角后再回来可重新探索。南峰亭和东峰亭每累计3次可再次触发。')
sub('第十二次的黑暗', '山洞第12次触发《大悲咒》时，所有人工灯光突然熄灭，包括山路灯、亭灯、灯塔和洞中转角微光；自然的日光或月光仍然保留。暂停、结束歌曲或退出触发视角后，灯光重新跟随环境恢复。想看清这处彩蛋，夜晚去山洞更有气氛。')
sub('把彩蛋留给一次散步', '可以先在全岛听一段《山歌好比春江水》，去码头等待《送别》，再沿山路看亭子、风铃与灯塔。最后进入山洞，一次次点击，直到原本温柔的海岛露出另一种气氛。也可以不按表寻找，只开着声音与自动轮换，等远海的船影、鲸鱼和下一次月夜自己出现。')

# Set explicit fonts and remove the default Word title border and theme overrides.
for node in doc.styles.element.iter():
    if node.tag == qn('w:rFonts'):
        for key in list(node.attrib):
            if 'theme' in key.lower(): del node.attrib[key]
    if node.tag == qn('w:pBdr'):
        node.getparent().remove(node)

doc.core_properties.title = '浮屿作品说明与彩蛋指南'
doc.core_properties.subject = '浮屿海岛作品的观察细节和音乐彩蛋'
doc.core_properties.author = ''
doc.core_properties.comments = ''
doc.save(OUT)
print(OUT)
